"""Formatted workbooks; all text is explicitly stored as text, never formulas."""
from collections import defaultdict
from io import BytesIO
from math import ceil

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def sheet(workbook, name, title, context, headers, rows, widths, filters=True):
    ws = workbook.create_sheet(name)
    for row in [[title], [context], headers, *rows]:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub('', cell.value)
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in (1, 3):
        for cell in ws[row]:
            cell.font = Font(bold=True, color='FFFFFF', size=11)
            cell.fill = PatternFill('solid', fgColor='153B46')
    for row in ws:
        lines = max(sum(max(1, ceil(len(line) / max(1, widths[cell.column-1]-2)))
                        for line in str(cell.value or '').split('\n')) for cell in row)
        ws.row_dimensions[row[0].row].height = min(409, max(24, lines * 15))
    ws.freeze_panes = 'B4' if name == 'Weekly Timetable' else 'A4'
    if filters:
        ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}'
    ws.print_title_rows = '1:3'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.dimensions
    ws.oddFooter.center.text = 'Page &P of &N'
    return ws


def schedule_workbook(title, term, entries, days):
    workbook = Workbook()
    workbook.remove(workbook.active)
    entries = sorted(entries, key=lambda e: (e.start_time, e.end_time, e.day, e.pk))
    slots = defaultdict(lambda: defaultdict(list))
    details = []
    for e in entries:
        course = e.effective_subject
        slots[(e.start_time, e.end_time)][e.day].append(
            f'{course.code} - {course.title}\n{e.effective_section}\n{e.effective_faculty}\n{e.room.name}\n{e.effective_component_display}')
        details.append([e.get_day_display(), e.start_time, e.end_time, str(course.code),
                        course.title, str(e.effective_section), str(e.effective_faculty),
                        e.room.name, e.effective_component_display])
    weekly = []
    for (start, end), groups in sorted(slots.items()):
        for index in range(max(map(len, groups.values()))):
            weekly.append([f'{start:%H:%M} - {end:%H:%M}'] +
                          [groups[day][index] if index < len(groups[day]) else '' for day, _ in days])
    sheet(workbook, 'Weekly Timetable', title, str(term), ['Time'] + [label for _, label in days],
          weekly or [['No scheduled meetings']], [18] + [28] * len(days), filters=False)
    ws = sheet(workbook, 'Schedule Details', title, str(term),
               ['Day', 'Start', 'End', 'Course Code', 'Course Title', 'Section', 'Faculty', 'Room', 'Session Type'],
               details, [14, 12, 12, 18, 38, 24, 28, 24, 18])
    for row in ws.iter_rows(min_row=4):
        row[1].number_format = row[2].number_format = 'hh:mm'
    return workbook


def workbook_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
