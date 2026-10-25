import io
import re
from datetime import time
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import SimpleTestCase
from django.urls import reverse
from reportlab.platypus import LongTable, Paragraph
from openpyxl import load_workbook

from . import models as m
from .services.pdf_export import write_table_pdf
from .test_known_defects import SchedulingDefectFixture


class PrintableTableTests(SimpleTestCase):
    def test_wrapping_and_repeated_headers(self):
        tables = []
        def capture(*args, **kwargs):
            table = LongTable(*args, **kwargs)
            tables.append(table)
            return table
        output = io.BytesIO()
        with patch('schedules.services.pdf_export.LongTable', side_effect=capture):
            write_table_pdf(output, 'Test', 'Term 2026', ['Course', 'Faculty'],
                            [['A < B & C ' * 30, 'Long faculty name ' * 12]], [2, 1])
        table = tables[0]
        self.assertEqual(table.repeatRows, 1)
        self.assertEqual(table.splitInRow, 1)
        self.assertIsInstance(table._cellvalues[1][0][0], Paragraph)
        self.assertGreater(table._rowHeights[1], 24)
        self.assertLessEqual(sum(table._colWidths), 732)
        self.assertTrue(output.getvalue().startswith(b'%PDF'))

    def test_many_entries_and_oversized_cell_paginate(self):
        output = io.BytesIO()
        write_table_pdf(output, 'Long report', 'Term', ['Course', 'Room'],
                        [['Long title ' * 1500, 'Room']] + [['Course', 'Room']] * 150, [2, 1])
        self.assertGreater(len(re.findall(rb'/Type /Page\b', output.getvalue())), 3)
        self.assertIn(b'/MediaBox [ 0 0 792 612 ]', output.getvalue())

    def test_empty_report_is_valid(self):
        output = io.BytesIO()
        write_table_pdf(output, 'Empty', 'Term', ['Course'], [], [1])
        self.assertTrue(output.getvalue().startswith(b'%PDF'))


class ScheduleExportTests(SchedulingDefectFixture):
    def workbook(self, **params):
        response = self.client.get(reverse('export_xlsx', args=[self.schedule.pk]), params)
        self.assertEqual(response.status_code, 200)
        return load_workbook(io.BytesIO(response.content))

    def setUp(self):
        self.admin = User.objects.create_superuser('export-admin')
        self.client.force_login(self.admin)
        self.schedule = m.Schedule.objects.create(term=self.term, name='Export example')
        for i, meeting in enumerate(self.meetings):
            self.schedule.entries.create(assignment=meeting, room=self.room, day=0,
                                         start_time=time(8+i), end_time=time(9+i))

    def test_pdf_and_csv_downloads(self):
        response = self.client.get(reverse('export_pdf', args=[self.schedule.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(b'/MediaBox [ 0 0 792 612 ]', response.content)
        self.assertEqual(self.client.get(reverse('export_csv', args=[self.schedule.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('room_utilization_pdf'), {'schedule': self.schedule.pk}).status_code, 200)

    def test_xlsx_contents_formats_filters_and_text_safety(self):
        m.Subject.objects.filter(pk=self.subject.pk).update(code='00123', title='=1+1')
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ['Weekly Timetable', 'Schedule Details'])
        details = workbook['Schedule Details']
        self.assertEqual(details.max_row, 5)
        self.assertEqual(details['D4'].value, '00123')
        self.assertEqual(details['E4'].value, '=1+1')
        self.assertEqual(details['E4'].data_type, 's')
        self.assertEqual(details['B4'].value, time(8))
        self.assertEqual(details['B4'].number_format, 'hh:mm')
        self.assertEqual(details['I4'].value, 'Lecture')
        self.assertEqual(details.freeze_panes, 'A4')
        self.assertEqual(details.auto_filter.ref, 'A3:I5')
        self.assertIn('00123', workbook['Weekly Timetable']['B4'].value)
        self.assertEqual(self.workbook(day=1)['Schedule Details'].max_row, 3)
        self.assertEqual(self.workbook(day='invalid')['Schedule Details'].max_row, 3)
        response = self.client.get(reverse('schedule_detail', args=[self.schedule.pk]))
        self.assertContains(response, reverse('export_xlsx', args=[self.schedule.pk]))

    def test_empty_and_room_exports(self):
        self.schedule.entries.all().delete()
        self.assertEqual(self.workbook()['Schedule Details'].max_row, 3)
        response = self.client.get(reverse('room_utilization_xlsx'), {'schedule': self.schedule.pk, 'q': 'does-not-match'})
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active.max_row, 3)

    def test_separate_session_types(self):
        m.Subject.objects.filter(pk=self.subject.pk).update(lecture_hours=1, lab_hours=1)
        m.TeachingAssignment.objects.filter(pk=self.meetings[1].pk).update(component='LABORATORY')
        details = self.workbook()['Schedule Details']
        self.assertEqual([details.cell(row, 9).value for row in (4, 5)], ['Lecture', 'Laboratory'])

    def test_role_visibility_and_published_snapshot(self):
        from .services.publication import publish
        user = User.objects.create_user('export-student')
        m.Profile.objects.create(user=user, role=m.Role.STUDENT)
        student = m.Student.objects.create(user=user, section=self.section, student_number='001', full_name='Example')
        self.client.force_login(user)
        url = reverse('export_xlsx', args=[self.schedule.pk])
        self.assertEqual(self.client.get(url).status_code, 404)
        other_section = m.Section.objects.create(program=self.section.program, year_level=self.section.year_level, name='Other', size=20)
        other_faculty = m.Faculty.objects.create(department=self.department, employee_id='EXPORT2', full_name='Private Faculty')
        parent = m.Assignment.objects.create(subject=self.subject, faculty=other_faculty, section=other_section, term=self.term)
        meeting = m.TeachingAssignment.objects.create(assignment=parent, component='LECTURE', duration_minutes=120, meeting_index=1)
        self.schedule.entries.create(assignment=meeting, room=self.room, day=0, start_time=time(10), end_time=time(12))
        faculty_user = User.objects.create_user('export-faculty')
        m.Profile.objects.create(user=faculty_user, role=m.Role.FACULTY)
        self.faculty.user = faculty_user
        self.faculty.save()
        publish(self.schedule, self.admin)
        self.assertEqual(self.workbook()['Schedule Details'].max_row, 5)
        self.assertEqual(self.client.get(reverse('room_utilization_xlsx')).status_code, 302)
        student.section = m.Section.objects.create(program=self.section.program, year_level=self.section.year_level, name='Unassigned', size=20)
        student.save()
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.force_login(faculty_user)
        details = self.workbook()['Schedule Details']
        self.assertEqual(details.max_row, 5)
        self.assertNotIn('Private Faculty', str(list(details.values)))
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)
